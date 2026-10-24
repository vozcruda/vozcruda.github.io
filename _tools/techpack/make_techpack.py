#!/usr/bin/env python3
"""Build the downloadable tech pack templates in /downloads.

  python3 _tools/techpack/make_techpack.py        (needs: pip install openpyxl)
  node _tools/techpack/make_pdf.js                (printable PDF, uses Playwright)

Outputs downloads/tech-pack-template.xlsx (blank), downloads/tech-pack-example-t-shirt.xlsx
(filled in with the example from /tech-pack-example-t-shirt.html) and _tools/techpack/print.html,
which make_pdf.js turns into downloads/tech-pack-template.pdf.
"""
import html
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "downloads"
SIZES = ["S", "M", "L", "XL"]

TEAL = PatternFill("solid", fgColor="1ABABD")
PALE = PatternFill("solid", fgColor="E6F7F7")
THIN = Side(style="thin", color="B7C2C3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# Each section: (sheet name, title, help text, columns, blank row count, example rows)
COVER = [("Brand", "Your brand name"), ("Style name", "Oversized Heavy Tee"), ("Style number", "TS-001"),
         ("Season or drop", "Drop 01"), ("Version and date", "V1, first issue"), ("Sizes", "S, M, L, XL"),
         ("Product type", "Oversized t-shirt"), ("Target quantity", "130 pieces"),
         ("Target delivery date", ""), ("Contact", "Name, phone and email of the person who approves samples")]
SECTIONS = [
    ("2 Sketches", "Flat sketches and callouts",
     "Insert front and back flat sketches (line drawings, lying flat, no model). List every detail the factory must not change.",
     ["Callout", "Detail"], 8,
     [["1", "Dropped shoulder seam"], ["2", "2 cm 1x1 rib neck, with self-fabric back neck tape"],
      ["3", "Single-needle topstitch at neck rib"], ["4", "2.5 cm twin-needle hem on body and sleeves"],
      ["5", "Front print position (see 6 Artwork) and printed neck label (see 7 Labels & packing)"]]),
    ("3 Size chart", "Size chart and tolerances (cm, garment measured flat)",
     "One row per point of measure. Add a small diagram with the letters so the quality checker measures the same way you do.",
     ["Point of measure"] + SIZES + ["Tolerance ±"], 12,
     [["A. Half chest, 2 cm below armhole", 56, 58.5, 61, 63.5, 1],
      ["B. Body length, from high point of shoulder", 70, 72, 74, 76, 1],
      ["C. Shoulder width, seam to seam", 54, 56, 58, 60, 1],
      ["D. Sleeve length, from shoulder seam", 22, 23, 24, 25, 0.5],
      ["E. Half sleeve opening", 20, 21, 22, 23, 0.5], ["F. Armhole, straight", 25, 26, 27, 28, 0.5],
      ["G. Neck width, seam to seam", 19, 19.5, 20, 20.5, 0.5], ["H. Front neck drop", 9, 9, 9.5, 9.5, 0.5],
      ["I. Neck rib width", 2, 2, 2, 2, 0.2], ["J. Half hem width", 56, 58.5, 61, 63.5, 1]]),
    ("4 BOM", "Bill of materials",
     "Every material and trim, with composition, GSM and finish. 'Cotton tee' could be quoted at 160 or 260 GSM.",
     ["Item", "Specification", "Colour", "Where used", "Supplier or reference"], 10,
     [["Main fabric", "100% combed cotton single jersey, 240 GSM, bio-washed, pre-shrunk", "Black, off-white", "Body and sleeves", ""],
      ["Rib", "1x1 cotton rib, matched to body colour", "Match body", "Neck", ""],
      ["Back neck tape", "Self fabric, 1 cm", "Match body", "Inside back neck", ""],
      ["Sewing thread", "Spun polyester", "Match body", "All seams", ""],
      ["Care label", "Printed satin, sewn into left side seam", "White", "Inside left side seam", ""]]),
    ("5 Construction", "Construction details",
     "Seams, stitch types, hem widths and finishes. These are easy to change silently between sample and bulk if not written down.",
     ["Area", "Instruction"], 8,
     [["Shoulder seams", "4-thread overlock with tape for stability"], ["Side seams", "4-thread overlock, no side slits"],
      ["Sleeve attach", "4-thread overlock"], ["Neck", "Rib attached with overlock, single-needle topstitch on body side"],
      ["Hems", "Body and sleeve hems 2.5 cm, twin-needle coverstitch"],
      ["Stitch density", "State your expectation, or ask the factory for its standard and record it here"]]),
    ("6 Artwork", "Artwork and placement",
     "One block per print or embroidery. Draw the placement measurement on the flat sketch too. Attach artwork as separate vector files.",
     ["Field", "Print 1", "Print 2"], 0,
     [["Method", "Screen print, water-based ink", ""], ["Artwork file", "TS-001_front_V1 (vector PDF, separate attachment)", ""],
      ["Print size", "28 cm wide x 20 cm high on size M; same size on all sizes", ""],
      ["Placement", "Centred, top of artwork 8 cm below front neck seam", ""],
      ["Colours", "One colour, off-white on black body (Pantone reference or swatch)", ""]]),
    ("7 Labels & packing", "Labels and packing",
     "Custom labels can have their own minimums and lead times, so order them early.",
     ["Item", "Instruction"], 0,
     [["Neck label", "Printed inside back neck, brand logo and size, 6 cm wide, centred 2 cm below neck seam"],
      ["Care label", "Fibre content, wash care symbols, country of origin, inside left side seam 8 cm above hem"],
      ["Hang tag", "Attached at left sleeve hem with a tag pin or string, if you use one"],
      ["Fold", "Folded with print facing out, size label visible"],
      ["Polybag", "One piece per polybag with size sticker"],
      ["Cartons", "Packed by size; carton marking shows style number, colour, size and quantity"]]),
]
ORDER_EXAMPLE = [["Black", 15, 25, 25, 15], ["Off-white", 10, 15, 15, 10]]
HELP = [
    "How to use this tech pack template",
    "",
    "A tech pack is the blueprint a factory uses to quote, sample and produce your garment.",
    "1. Fill in 1 Cover first. The style number and version appear at the top of every sheet automatically.",
    "2. Work through each sheet. Delete example rows you don't need and add rows as required.",
    "3. Insert your flat sketches and artwork images on 2 Sketches and 6 Artwork, and also send artwork as separate vector files.",
    "4. When a sample comes back, review it against this file point by point, then issue V2 with the changes.",
    "",
    "Guides: https://vozcruda.com/tech-pack-guide-for-beginners.html",
    "Worked example: https://vozcruda.com/tech-pack-example-t-shirt.html",
    "Order below a factory's minimum: https://smallotz.vozcruda.com/",
    "",
    "Free template by Voz Cruda (vozcruda.com). You may use and adapt it for your own products.",
]

def header_block(ws, title, help_text, width):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Style no.:"
    ws["B2"] = "='1 Cover'!B5"
    ws["C2"] = "Version:"
    ws["D2"] = "='1 Cover'!B7"
    for c in ("A2", "C2"):
        ws[c].font = Font(bold=True, color="5A6668")
    ws["A3"] = help_text
    ws["A3"].font = Font(italic=True, color="5A6668")
    ws["A3"].alignment = WRAP
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(width, 4))
    ws.row_dimensions[3].height = 32

def table(ws, row, cols, rows, widths):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border, cell.alignment = Font(bold=True, color="0A0A0A"), TEAL, BOX, WRAP
    for r, vals in enumerate(rows, row + 1):
        for i in range(1, len(cols) + 1):
            v = vals[i - 1] if i - 1 < len(vals) else None
            cell = ws.cell(row=r, column=i, value=v if v != "" else None)
            cell.border, cell.alignment = BOX, WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def build(example):
    wb = Workbook()
    ws = wb.active
    ws.title = "How to use"
    for i, line in enumerate(HELP, 1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110

    ws = wb.create_sheet("1 Cover")
    ws["A1"] = "Tech pack: cover"
    ws["A1"].font = Font(bold=True, size=14)
    table(ws, 2, ["Field", "Details"], [[k, v if example else ""] for k, v in COVER], [26, 70])
    # B5 = style number, B7 = version (referenced by every other sheet)

    for name, title, help_text, cols, blank, rows in SECTIONS:
        ws = wb.create_sheet(name)
        header_block(ws, title, help_text, len(cols))
        if example:
            data = rows + [[""] * len(cols)] * 3
        elif blank:
            data = [[""] * len(cols)] * blank
        else:  # field-style tables keep their labels
            data = [[r[0]] + [""] * (len(cols) - 1) for r in rows] + [[""] * len(cols)] * 2
        if name == "3 Size chart" and not example:
            data = [[r[0]] + [""] * (len(cols) - 1) for r in rows] + [[""] * len(cols)] * 2
        widths = {"3 Size chart": [44, 9, 9, 9, 9, 13], "4 BOM": [18, 52, 18, 22, 24],
                  "6 Artwork": [18, 60, 40]}.get(name, [22, 90])
        table(ws, 5, cols, data, widths)

    ws = wb.create_sheet("8 Order sheet")
    header_block(ws, "Order sheet", "Quantity per colour and size. Totals calculate automatically.", 6)
    rows = ORDER_EXAMPLE if example else [[""] * 5] * 4
    rows = [list(r) for r in rows] + ([[""] * 5] * 2 if example else [])
    table(ws, 5, ["Colour"] + SIZES + ["Total"], rows, [22, 9, 9, 9, 9, 11])
    first, last = 6, 5 + len(rows)
    for r in range(first, last + 1):
        ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})").font = Font(bold=True)
    tr = last + 1
    ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
    for c in range(2, 7):
        L = get_column_letter(c)
        cell = ws.cell(row=tr, column=c, value=f"=SUM({L}{first}:{L}{last})")
        cell.font = Font(bold=True)
    for c in range(1, 7):
        ws.cell(row=tr, column=c).fill, ws.cell(row=tr, column=c).border = PALE, BOX
    ws.cell(row=tr + 2, column=1, value="Target delivery date:").font = Font(bold=True)
    ws.cell(row=tr + 3, column=1, value="Delivery address:").font = Font(bold=True)

    for s in wb.worksheets:
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.fitToHeight = 0
    return wb

def print_html():
    """Blank printable version: one A4 landscape page per section."""
    e = html.escape
    def tbl(cols, rows):
        head = "".join(f"<th>{e(c)}</th>" for c in cols)
        body = "".join("<tr>" + "".join(f"<td>{e(str(v))}</td>" for v in r) + "</tr>" for r in rows)
        return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"
    def page(n, title, help_text, inner):
        return (f'<section><header><div><b>{n}. {e(title)}</b><p>{e(help_text)}</p></div>'
                f'<div class="meta">Style no. ________ &nbsp; Version ____ &nbsp; Page {n}/8</div></header>{inner}'
                f'<footer>Free tech pack template · vozcruda.com/tech-pack-template.html</footer></section>')
    pages = [page(1, "Cover", "Fill this page first. Use the same style number and version on every page.",
                  tbl(["Field", "Details"], [[k, ""] for k, _ in COVER]))]
    for n, (name, title, help_text, cols, blank, rows) in enumerate(SECTIONS, 2):
        if name == "2 Sketches":
            inner = '<div class="sketch"><div>Front</div><div>Back</div></div>' + tbl(cols, [["", ""]] * 6)
        elif blank and name != "3 Size chart":
            inner = tbl(cols, [[""] * len(cols)] * blank)
        else:
            inner = tbl(cols, [[r[0]] + [""] * (len(cols) - 1) for r in rows] + [[""] * len(cols)] * 2)
        pages.append(page(n, title, help_text, inner))
    pages.append(page(8, "Order sheet", "Quantity per colour and size. Add the target delivery date and address.",
                      tbl(["Colour"] + SIZES + ["Total"], [[""] * 6] * 6 + [["Total"] + [""] * 5])
                      + '<p class="line">Target delivery date: ____________________ &nbsp; Delivery address: ______________________________________________</p>'))
    css = """@page{size:A4 landscape;margin:12mm}
body{font:10.5pt/1.4 Helvetica,Arial,sans-serif;color:#111;margin:0}
section{page-break-after:always;display:flex;flex-direction:column;min-height:180mm}
section:last-child{page-break-after:auto}
header{display:flex;justify-content:space-between;gap:20px;border-bottom:2px solid #1ABABD;padding-bottom:6px;margin-bottom:10px}
header b{font-size:15pt}header p{margin:3px 0 0;color:#555;max-width:170mm}
.meta{font-size:9pt;color:#333;white-space:nowrap;align-self:flex-end}
table{border-collapse:collapse;width:100%}th,td{border:1px solid #9aa5a6;padding:6px 8px;text-align:left;vertical-align:top;height:16px}
th{background:#E6F7F7;font-size:9pt;text-transform:uppercase;letter-spacing:.04em}
.sketch{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin-bottom:10px}.sketch div{border:1px dashed #9aa5a6;height:78mm;padding:6px;color:#777}
footer{margin-top:auto;padding-top:8px;font-size:8pt;color:#777}.line{margin-top:14px}"""
    return f"<!doctype html><html><head><meta charset='utf-8'><title>Tech pack template</title><style>{css}</style></head><body>{''.join(pages)}</body></html>"

def main():
    OUT.mkdir(exist_ok=True)
    build(False).save(OUT / "tech-pack-template.xlsx")
    build(True).save(OUT / "tech-pack-example-t-shirt.xlsx")
    (Path(__file__).parent / "print.html").write_text(print_html(), encoding="utf-8")
    print("Wrote downloads/tech-pack-template.xlsx, downloads/tech-pack-example-t-shirt.xlsx, _tools/techpack/print.html")

if __name__ == "__main__":
    main()
